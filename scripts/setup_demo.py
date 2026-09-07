#!/usr/bin/env python3
"""Build a wholly synthetic, offline demo and configure its Power BI sources."""

import argparse
import json
import re
from datetime import date, datetime, timedelta
from pathlib import Path

from openpyxl import Workbook

ROOT = Path(__file__).resolve().parents[1]
MODEL = ROOT / 'pbip/Psiko Demo.SemanticModel/definition'
PROJECT = 'pbip/Psiko Demo.pbip'
MARKER = "demo-manifest.json"


def write_book(folder, filename, sheet, columns, rows, date_columns=()):
    """Write sample rows with the exact workbook/sheet names used by the model."""
    workbook = Workbook()
    workbook.properties.creator = "Synthetic portfolio demo"
    workbook.properties.description = "Invented sample data; not business results."
    worksheet = workbook.active
    worksheet.title = sheet
    worksheet.append(columns)
    for row in rows:
        if len(row) != len(columns):
            raise ValueError(f"Wrong column count in {filename}")
        worksheet.append(row)
    for column in date_columns:
        index = columns.index(column) + 1
        for row in range(2, worksheet.max_row + 1):
            worksheet.cell(row, index).number_format = "yyyy-mm-dd"
    worksheet.freeze_panes = "A2"
    worksheet.auto_filter.ref = worksheet.dimensions
    workbook.save(folder / filename)
    workbook.close()
    return {"file": filename, "sheet": sheet, "rows": len(rows), "columns": columns}


def configure_model(model, folder):
    expressions = model / "expressions.tmdl"
    source = expressions.read_text(encoding="utf-8")
    # Power Query text literals escape quotes by doubling them.
    value = folder.resolve().as_posix().replace('"', '""')
    replacement = ('expression DemoDataFolder = "' + value
                   + '" meta [IsParameterQuery=true, Type="Text", IsParameterQueryRequired=true]')
    source, count = re.subn(r'^expression DemoDataFolder = .*$',
                            lambda match: replacement, source, flags=re.MULTILINE)
    if count != 1:
        raise ValueError("Expected one DemoDataFolder parameter in expressions.tmdl")
    expressions.write_text(source, encoding="utf-8")


def build_data(folder):
    trainee_columns = ["ADAY NO", "SERTİFİKA", "SERTİFİKA SERİ NO", "SERTİFİKA VER.TARİHİ",
                       "TELEFON", "Candidate", "İKİNCİ/DİREK."]
    result_columns = ["ADAY NO", "ADI", "SOYADI", "SERT.VER.TARİHİ", "P.TEKNİK ALIŞ TARİHİ",
                      "Ay", "Yıl", "SERTİFİKA", "TELEFON", "Sıradaki Test", "Durum"]
    trainees, results = [], []
    for i in range(1, 37):
        candidate = 2000 + i
        issued = date(2019 + (i - 1) % 7, 1 + (i - 1) % 12, 15)
        licence = ["C", "D", "E"][(i - 1) % 3]
        phone = f"DEMO-PHONE-{i:03}"
        trainees.append([candidate, licence, f"DEMO-SERIAL-{i:03}", issued,
                         phone, f"DEMO CANDIDATE {i:03}", "DİREK"])
        # Illustrate interested, not interested, and not called states.
        assessment = [date(2021 + (i - 1) % 5, 1 + (i - 1) % 12, 15), date(1968, 1, 1), None][i % 3]
        results.append([candidate, "DEMO", f"CANDIDATE {i:03}", issued, assessment,
                        issued.strftime("%B"), issued.year, licence, phone,
                        date(2026, 6, 30), "Demo"])
    pool_columns = ["SN", "Customer", "TARİH", "TELEFON"]
    pool_2025 = [[3000 + i, f"DEMO EXTERNAL 2025-{i:02}", date(2025, i, 15), f"DEMO-2025-{i:02}"]
                 for i in range(1, 13)]
    pool_2026 = [[4000 + i, f"DEMO EXTERNAL 2026-{i:02}", date(2026, i, 15), f"DEMO-2026-{i:02}"]
                 for i in range(1, 13)]
    return [
        write_book(folder, "2010-2025 TARAMA.xlsx", "Sayfa1", trainee_columns, trainees,
                   ["SERTİFİKA VER.TARİHİ"]),
        write_book(folder, "22.06.2026-son2507.xlsx", "Export", result_columns, results,
                   ["SERT.VER.TARİHİ", "P.TEKNİK ALIŞ TARİHİ", "Sıradaki Test"]),
        write_book(folder, "BI-Psiko 2025.xlsx", "TEMMUZ", pool_columns, pool_2025, ["TARİH"]),
        write_book(folder, "BI-Psiko 2026.xlsx", "2026", pool_columns, pool_2026, ["TARİH"]),
    ]



def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--data-dir", type=Path, default=ROOT / "demo-data",
                        help="Output folder for generated workbooks (default: repo/demo-data)")
    args = parser.parse_args()
    folder = args.data_dir.expanduser().resolve()
    marker = folder / MARKER
    if folder.exists() and any(folder.glob("*.xlsx")):
        if not marker.exists():
            parser.error("Output contains existing workbooks; choose a new --data-dir.")
        if json.loads(marker.read_text(encoding="utf-8")).get("generator") != PROJECT:
            parser.error("Output belongs to a different demo; choose a new --data-dir.")
    folder.mkdir(parents=True, exist_ok=True)
    files = build_data(folder)
    marker.write_text(json.dumps({
        "generator": PROJECT,
        "data_kind": "fully synthetic; no private inputs or API calls",
        "reference_year": 2026,
        "files": files,
    }, ensure_ascii=False, indent=2) + "\n", encoding="utf-8")
    configure_model(MODEL, folder)
    print(f"Created {len(files)} synthetic workbooks in {folder}")
    print(f"Configured DemoDataFolder. Open {ROOT / PROJECT} in Power BI Desktop and Refresh.")


if __name__ == "__main__":
    main()
