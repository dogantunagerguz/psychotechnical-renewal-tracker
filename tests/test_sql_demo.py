import sys
from contextlib import contextmanager
from datetime import date
from pathlib import Path
from tempfile import TemporaryDirectory
import unittest

from openpyxl import load_workbook

ROOT = Path(__file__).resolve().parents[1]
sys.path.insert(0, str(ROOT / "sql"))
import run_demo  # noqa: E402


class SqlPortfolioDemoTests(unittest.TestCase):
    def setUp(self):
        self.connection = run_demo.build_connection()

    def tearDown(self):
        self.connection.close()

    @contextmanager
    def mutation(self):
        self.connection.execute("SAVEPOINT mutation")
        try:
            yield
        finally:
            self.connection.execute("ROLLBACK TO mutation")
            self.connection.execute("RELEASE mutation")

    def issues(self):
        return dict(self.connection.execute("SELECT * FROM quality_check_results"))

    def kpis(self):
        return dict(self.connection.execute("SELECT * FROM contact_outcome_kpis").fetchone())

    def assert_not_actionable(self, candidate_id):
        self.assertIsNone(self.connection.execute(
            "SELECT 1 FROM call_queue_2026 WHERE candidate_id = ?", (candidate_id,)
        ).fetchone())
        self.assertIsNone(self.connection.execute(
            "SELECT 1 FROM unified_contact_pool WHERE pool_key = ?",
            (f"TRAINEE-{candidate_id}",)
        ).fetchone())

    def test_pipeline_loads_all_public_demo_sources(self):
        row = self.connection.execute("SELECT * FROM pipeline_reconciliation").fetchone()
        self.assertEqual(tuple(row), (36, 36, 24, 20, 4, 36, 24, 32))

    def test_two_regime_anchor_and_due_date(self):
        row = self.connection.execute("""
            SELECT anchor_date, next_test_date, renewal_status
            FROM renewal_priority_2026
            WHERE candidate_id = 2001
        """).fetchone()
        self.assertEqual(tuple(row), ("2021-06-30", "2026-06-30", "Overdue"))

    def test_contact_rate_uses_called_denominator(self):
        row = self.connection.execute("SELECT * FROM contact_outcome_kpis").fetchone()
        self.assertEqual(tuple(row), (36, 24, 12, 12, 50.0, 12, 0, 0))

    def test_invalid_assessments_are_preserved_quarantined_and_not_interested(self):
        bad_values = ["", "   ", "\t", "not-a-date", "2026-02-30", "2025-02-29",
                      "2026-00-01", "2026-13-01", "2026-2-03", "2026-09-09",
                      "2026-01-01 00:00:00", "0000-01-01", "2026-01-00"]
        for value in bad_values:
            with self.subTest(value=value), self.mutation():
                self.connection.execute(
                    "UPDATE raw_call_results SET psychotechnical_assessment_date = ? WHERE candidate_id = 2003",
                    (value,)
                )
                raw = self.connection.execute(
                    "SELECT psychotechnical_assessment_date FROM raw_call_results WHERE candidate_id = 2003"
                ).fetchone()[0]
                self.assertEqual(raw, value)
                self.assertEqual(self.connection.execute(
                    "SELECT contact_status FROM stg_call_outcomes WHERE candidate_id = 2003"
                ).fetchone()[0], "Quality error")
                self.assertEqual(self.issues()["invalid_or_future_assessment_date"], 1)
                self.assertEqual(self.kpis()["interested"], 11)
                self.assertEqual(self.kpis()["called_trainees"], 23)
                self.assertEqual(self.kpis()["quality_error_trainees"], 1)
                self.assert_not_actionable(2003)

    def test_null_sentinel_valid_leap_day_and_as_of_date_have_distinct_meanings(self):
        for value, expected in [(None, "Not called"), ("1968-01-01", "Not interested"),
                                ("2024-02-29", "Interested"), ("2026-09-08", "Interested")]:
            with self.subTest(value=value), self.mutation():
                self.connection.execute(
                    "UPDATE raw_call_results SET psychotechnical_assessment_date = ? WHERE candidate_id = 2002",
                    (value,)
                )
                self.assertEqual(self.connection.execute(
                    "SELECT contact_status FROM stg_call_outcomes WHERE candidate_id = 2002"
                ).fetchone()[0], expected)
                self.assertEqual(self.issues()["invalid_or_future_assessment_date"], 0)

    def test_missing_result_is_quality_error_not_not_called(self):
        self.connection.execute("DELETE FROM raw_call_results WHERE candidate_id = 2002")
        self.assertEqual(self.issues()["trainee_without_result"], 1)
        self.assertEqual(self.kpis()["eligible_trainees"], 36)
        self.assertEqual(self.kpis()["not_called_trainees"], 11)
        self.assertEqual(self.kpis()["quality_error_trainees"], 1)
        self.assert_not_actionable(2002)

    def test_invalid_identity_values_are_not_cast_into_valid_people(self):
        for value in [None, "", " ", "bad", "2003x", "02003", "2003", 0, -1, 2003.5]:
            with self.subTest(value=value), self.mutation():
                self.connection.execute(
                    "UPDATE raw_trainees SET candidate_id = ? WHERE candidate_id = 2003", (value,)
                )
                self.assertEqual(self.issues()["invalid_trainee_id"], 1)
                self.assertEqual(self.kpis()["excluded_identity_rows"], 1)
                self.assertEqual(self.kpis()["eligible_trainees"], 35)
                self.assert_not_actionable(2003)
                self.assertEqual(self.connection.execute(
                    "SELECT candidate_id FROM raw_trainees WHERE candidate_name = 'DEMO CANDIDATE 003'"
                ).fetchone()[0], value)

    def test_invalid_result_and_external_identity_are_visible(self):
        self.connection.execute("UPDATE raw_call_results SET candidate_id = '2003x' WHERE candidate_id = 2003")
        self.connection.execute("UPDATE raw_external_pool SET pool_id = NULL WHERE pool_id = 3001")
        checks = self.issues()
        self.assertEqual(checks["invalid_result_id"], 1)
        self.assertEqual(checks["trainee_without_result"], 1)
        self.assertEqual(checks["invalid_external_pool_id"], 1)
        self.assertEqual(self.kpis()["quality_error_trainees"], 1)
        self.assert_not_actionable(2003)
        self.assertEqual(self.connection.execute(
            "SELECT COUNT(*) FROM unified_contact_pool WHERE source = 'External'"
        ).fetchone()[0], 19)

    def test_missing_contact_references_block_actionable_rows(self):
        for table, key, value, check in [
            ("raw_trainees", "candidate_id", 2002, "missing_trainee_contact_reference"),
            ("raw_call_results", "candidate_id", 2002, "missing_result_contact_reference"),
            ("raw_external_pool", "pool_id", 3001, "missing_external_contact_reference"),
        ]:
            for blank in (None, "", " \t\r\n "):
                with self.subTest(table=table, blank=blank), self.mutation():
                    self.connection.execute(f"UPDATE {table} SET phone_ref = ? WHERE {key} = ?", (blank, value))
                    self.assertEqual(self.issues()[check], 1)
                    if table == "raw_external_pool":
                        self.assertEqual(self.connection.execute(
                            "SELECT COUNT(*) FROM unified_contact_pool WHERE source = 'External'"
                        ).fetchone()[0], 19)
                    else:
                        self.assert_not_actionable(2002)
                        self.assertEqual(self.kpis()["quality_error_trainees"], 1)

    def test_invalid_source_dates_do_not_create_schedule_or_pool_entries(self):
        for table, column, key, value, check in [
            ("raw_trainees", "licence_issue_date", "candidate_id", 2002, "trainee_row_quality_error"),
            ("raw_call_results", "certificate_issue_date", "candidate_id", 2002, "result_row_quality_error"),
            ("raw_external_pool", "source_date", "pool_id", 3001, "external_pool_row_quality_error"),
        ]:
            with self.subTest(table=table), self.mutation():
                self.connection.execute(f"UPDATE {table} SET {column} = '2026-02-30' WHERE {key} = ?", (value,))
                self.assertEqual(self.issues()[check], 1)
                if table != "raw_external_pool":
                    self.assert_not_actionable(2002)
                else:
                    self.assertEqual(self.connection.execute(
                        "SELECT COUNT(*) FROM unified_contact_pool WHERE source = 'External'"
                    ).fetchone()[0], 19)

    def test_duplicate_trainee_and_conflicting_result_do_not_multiply_outputs(self):
        for table, check in [("raw_trainees", "duplicate_trainee_id"),
                              ("raw_call_results", "duplicate_result_candidate")]:
            with self.subTest(table=table), self.mutation():
                self.connection.execute(f"INSERT INTO {table} SELECT * FROM {table} WHERE candidate_id = 2003")
                if table == "raw_call_results":
                    self.connection.execute("""
                        UPDATE raw_call_results SET psychotechnical_assessment_date = '1968-01-01'
                        WHERE rowid = (SELECT MAX(rowid) FROM raw_call_results)
                    """)
                    self.assertEqual(self.kpis()["quality_error_trainees"], 1)
                else:
                    self.assertEqual(self.kpis()["excluded_identity_rows"], 2)
                self.assertEqual(self.issues()[check], 1)
                self.assertEqual(self.kpis()["called_trainees"], 23)
                self.assert_not_actionable(2003)

    def test_conflicting_source_attributes_are_flagged_and_not_actionable(self):
        self.connection.execute("UPDATE raw_call_results SET licence_class = 'C' WHERE candidate_id = 2002")
        self.assertEqual(self.issues()["result_trainee_source_mismatch"], 1)
        self.assertEqual(self.kpis()["quality_error_trainees"], 1)
        self.assert_not_actionable(2002)

    def test_external_identity_is_workbook_and_id_and_duplicates_are_quarantined(self):
        self.connection.execute("UPDATE raw_external_pool SET pool_id = 3001 WHERE pool_id = 4001")
        self.assertEqual(self.issues()["external_pool_duplicate_key"], 0)
        keys = [row[0] for row in self.connection.execute(
            "SELECT pool_key FROM unified_contact_pool WHERE source = 'External'"
        )]
        self.assertEqual(len(keys), 20)
        self.assertEqual(len(set(keys)), 20)
        self.connection.execute("""
            INSERT INTO raw_external_pool SELECT * FROM raw_external_pool
            WHERE pool_id = 3001 AND source_workbook = 'BI-Psiko 2025.xlsx'
        """)
        self.assertEqual(self.issues()["external_pool_duplicate_key"], 1)
        self.assertEqual(self.connection.execute(
            "SELECT COUNT(*) FROM unified_contact_pool WHERE source = 'External'"
        ).fetchone()[0], 19)

    def test_zero_called_and_empty_population_have_zero_counts_and_null_rate(self):
        self.connection.execute("UPDATE raw_call_results SET psychotechnical_assessment_date = NULL")
        self.assertEqual(self.kpis()["called_trainees"], 0)
        self.assertEqual(self.kpis()["interest_rate_among_called_pct"], None)
        self.assertEqual(self.kpis()["not_called_trainees"], 36)
        for table in ("raw_trainees", "raw_call_results", "raw_external_pool"):
            self.connection.execute(f"DELETE FROM {table}")
        self.assertEqual(tuple(self.connection.execute("SELECT * FROM contact_outcome_kpis").fetchone()),
                         (0, 0, 0, 0, None, 0, 0, 0))

    def test_unified_pool_preserves_source(self):
        rows = self.connection.execute("""
            SELECT source, COUNT(*) FROM unified_contact_pool GROUP BY source
        """).fetchall()
        self.assertEqual(dict(rows), {"External": 20, "Interested trainee": 12})

    def test_all_quality_checks_pass(self):
        failures = self.connection.execute(
            "SELECT check_name FROM quality_check_results WHERE issue_count <> 0"
        ).fetchall()
        self.assertEqual(failures, [])

    def test_committed_results_snapshot_is_current(self):
        expected = run_demo.render_results(self.connection)
        actual = (ROOT / "sql/RESULTS.md").read_text(encoding="utf-8")
        self.assertEqual(actual, expected)

    def database_snapshot(self):
        queries = {
            "trainees": "SELECT * FROM raw_trainees ORDER BY rowid",
            "results": "SELECT * FROM raw_call_results ORDER BY rowid",
            "external": "SELECT * FROM raw_external_pool ORDER BY rowid",
            "queue": "SELECT * FROM call_queue_2026 ORDER BY candidate_id",
            "pool": "SELECT * FROM unified_contact_pool ORDER BY pool_key",
            "kpis": "SELECT * FROM contact_outcome_kpis",
        }
        return {name: [tuple(row) for row in self.connection.execute(query)]
                for name, query in queries.items()}

    def test_snapshot_replay_on_same_connection_preserves_values_and_outputs(self):
        with TemporaryDirectory() as directory:
            folder = Path(directory)
            run_demo.load_setup_module().build_data(folder)
            before = self.database_snapshot()
            run_demo.load_workbooks(self.connection, folder)
            self.assertEqual(self.database_snapshot(), before)
            run_demo.load_workbooks(self.connection, folder)
            self.assertEqual(self.database_snapshot(), before)
            self.assertTrue(all(count == 0 for count in self.issues().values()))

    def test_changed_snapshot_add_update_delete_and_replay_are_applied_once(self):
        with TemporaryDirectory() as directory:
            folder = Path(directory)
            run_demo.load_setup_module().build_data(folder)

            # Remove trainee 2001 and add a new, uncalled trainee based on 2002.
            path = folder / "2010-2025 TARAMA.xlsx"
            book = load_workbook(path)
            sheet = book.active
            new_trainee = [cell.value for cell in sheet[3]]
            new_trainee[0], new_trainee[5] = 2099, "DEMO NEW CANDIDATE"
            sheet.delete_rows(2)
            sheet.append(new_trainee)
            book.save(path)
            book.close()

            path = folder / "22.06.2026-son2507.xlsx"
            book = load_workbook(path)
            sheet = book.active
            new_result = [cell.value for cell in sheet[3]]
            new_result[0], new_result[2] = 2099, "NEW CANDIDATE"
            # 2002 moves from Not called to Interested in snapshot B.
            sheet.cell(3, 5, date(2026, 9, 8))
            sheet.delete_rows(2)
            sheet.append(new_result)
            book.save(path)
            book.close()

            # Also remove an external source record: full replacement must remove it.
            path = folder / "BI-Psiko 2025.xlsx"
            book = load_workbook(path)
            book.active.delete_rows(2)
            book.save(path)
            book.close()

            run_demo.load_workbooks(self.connection, folder)
            self.assertIsNone(self.connection.execute(
                "SELECT 1 FROM raw_trainees WHERE candidate_id = 2001"
            ).fetchone())
            self.assertIsNotNone(self.connection.execute(
                "SELECT 1 FROM call_queue_2026 WHERE candidate_id = 2099"
            ).fetchone())
            self.assertIsNone(self.connection.execute(
                "SELECT 1 FROM call_queue_2026 WHERE candidate_id = 2002"
            ).fetchone())
            self.assertIsNotNone(self.connection.execute(
                "SELECT 1 FROM unified_contact_pool WHERE pool_key = 'TRAINEE-2002'"
            ).fetchone())
            self.assertEqual(self.kpis()["interested"], 13)
            self.assertEqual(self.kpis()["not_interested"], 11)
            self.assertEqual(self.kpis()["called_trainees"], 24)
            self.assertEqual(self.kpis()["eligible_trainees"], 36)
            after = self.database_snapshot()
            run_demo.load_workbooks(self.connection, folder)
            self.assertEqual(self.database_snapshot(), after)
            self.assertTrue(all(count == 0 for count in self.issues().values()))

    def test_late_workbook_failure_rolls_back_all_source_tables(self):
        with TemporaryDirectory() as directory:
            folder = Path(directory)
            run_demo.load_setup_module().build_data(folder)
            # An earlier workbook would change the committed candidate's name.
            path = folder / "2010-2025 TARAMA.xlsx"
            book = load_workbook(path)
            book.active.cell(2, 6, "SHOULD NOT BE COMMITTED")
            book.save(path)
            book.close()
            # The final workbook fails only after the previous sources were inserted.
            (folder / "BI-Psiko 2026.xlsx").unlink()
            before = self.database_snapshot()
            with self.assertRaises(FileNotFoundError):
                run_demo.load_workbooks(self.connection, folder)
            self.assertEqual(self.database_snapshot(), before)

    def test_invalid_workbook_cell_is_retained_for_diagnosis(self):
        with TemporaryDirectory() as directory:
            folder = Path(directory)
            run_demo.load_setup_module().build_data(folder)
            path = folder / "22.06.2026-son2507.xlsx"
            book = load_workbook(path)
            book.active.cell(4, 5, "2026-02-30")
            book.save(path)
            book.close()
            run_demo.load_workbooks(self.connection, folder)
            self.assertEqual(self.connection.execute("""
                SELECT psychotechnical_assessment_date FROM raw_call_results WHERE candidate_id = 2003
            """).fetchone()[0], "2026-02-30")
            self.assertEqual(self.issues()["invalid_or_future_assessment_date"], 1)
            self.assert_not_actionable(2003)

    def test_bad_headers_even_without_rows_roll_back_snapshot(self):
        for malformed in ("header_only", "duplicate_header"):
            with self.subTest(malformed=malformed), TemporaryDirectory() as directory:
                folder = Path(directory)
                run_demo.load_setup_module().build_data(folder)
                path = folder / "BI-Psiko 2026.xlsx"
                book = load_workbook(path)
                if malformed == "header_only":
                    book.active.delete_rows(2, book.active.max_row)
                    book.active.cell(1, 1, "wrong header")
                else:
                    book.active.cell(1, 2, "SN")
                book.save(path)
                book.close()
                before = self.database_snapshot()
                with self.assertRaisesRegex(ValueError, "Invalid or ambiguous workbook headers"):
                    run_demo.load_workbooks(self.connection, folder)
                self.assertEqual(self.database_snapshot(), before)

    def test_uncached_assessment_formula_is_preserved_not_treated_as_not_called(self):
        with TemporaryDirectory() as directory:
            folder = Path(directory)
            run_demo.load_setup_module().build_data(folder)
            path = folder / "22.06.2026-son2507.xlsx"
            book = load_workbook(path)
            formula = "=DATE(2026,2,30)"
            book.active.cell(3, 5, formula)
            book.save(path)
            book.close()
            run_demo.load_workbooks(self.connection, folder)
            self.assertEqual(self.connection.execute("""
                SELECT psychotechnical_assessment_date FROM raw_call_results WHERE candidate_id = 2002
            """).fetchone()[0], formula)
            self.assertEqual(self.issues()["invalid_or_future_assessment_date"], 1)
            self.assertEqual(self.kpis()["quality_error_trainees"], 1)
            self.assert_not_actionable(2002)


if __name__ == "__main__":
    unittest.main()
