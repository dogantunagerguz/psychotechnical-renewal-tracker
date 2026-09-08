#!/usr/bin/env python3
"""Load the Power BI demo workbooks, execute SQL, and render reviewable results."""

import argparse
from datetime import date, datetime
import importlib.util
from pathlib import Path
import sqlite3
from tempfile import TemporaryDirectory

from openpyxl import load_workbook

SQL_DIR = Path(__file__).resolve().parent
ROOT = SQL_DIR.parent
SQL_FILES = [
    "01_raw_schema.sql",
    "02_staging.sql",
    "03_marts.sql",
    "04_quality_checks.sql",
]
REPORTS = [
    ("Pipeline reconciliation", "SELECT * FROM pipeline_reconciliation"),
    ("Renewal-status distribution", """
        SELECT renewal_status, COUNT(*) AS trainees
        FROM renewal_priority_2026
        GROUP BY renewal_status
        ORDER BY CASE renewal_status
            WHEN 'Overdue' THEN 1 WHEN 'Due now' THEN 2
            WHEN 'Upcoming' THEN 3 ELSE 4 END
    """),
    ("Open call queue sample", """
        SELECT candidate_name, licence_class, next_test_date, renewal_status
        FROM call_queue_2026
        ORDER BY priority_order, next_test_date, candidate_name
        LIMIT 10
    """),
    ("Synthetic contact KPI", "SELECT * FROM contact_outcome_kpis"),
    ("Unified-pool sources", """
        SELECT source, COUNT(*) AS records
        FROM unified_contact_pool
        GROUP BY source ORDER BY source
    """),
    ("Recent external-pool source trend", """
        SELECT * FROM external_pool_monthly_trend
        WHERE source_month >= '2025-10'
        ORDER BY source_month
    """),
    ("Data-quality checks", "SELECT * FROM quality_check_results ORDER BY check_name"),
]


def load_setup_module():
    path = ROOT / "scripts/setup_demo.py"
    spec = importlib.util.spec_from_file_location("psychotechnical_setup_demo", path)
    module = importlib.util.module_from_spec(spec)
    spec.loader.exec_module(module)
    return module


def rows_from_workbook(path):
    workbook = load_workbook(path, data_only=True, read_only=True)
    try:
        worksheet = workbook.active
        values = worksheet.iter_rows(values_only=True)
        headers = next(values)
        return [dict(zip(headers, row)) for row in values]
    finally:
        workbook.close()


def iso_date(value):
    if isinstance(value, datetime):
        return value.date().isoformat()
    if isinstance(value, date):
        return value.isoformat()
    return value


def load_workbooks(connection, folder):
    trainees = rows_from_workbook(folder / "2010-2025 TARAMA.xlsx")
    connection.executemany(
        "INSERT INTO raw_trainees VALUES (?, ?, ?, ?, ?, ?, ?)",
        [(
            row["ADAY NO"], row["SERTİFİKA"], row["SERTİFİKA SERİ NO"],
            iso_date(row["SERTİFİKA VER.TARİHİ"]), row["TELEFON"],
            row["Candidate"], row["İKİNCİ/DİREK."]
        ) for row in trainees],
    )

    results = rows_from_workbook(folder / "22.06.2026-son2507.xlsx")
    connection.executemany(
        "INSERT INTO raw_call_results VALUES (?, ?, ?, ?, ?, ?, ?, ?, ?, ?, ?)",
        [(
            row["ADAY NO"], row["ADI"], row["SOYADI"], iso_date(row["SERT.VER.TARİHİ"]),
            iso_date(row["P.TEKNİK ALIŞ TARİHİ"]), row["Ay"], row["Yıl"],
            row["SERTİFİKA"], row["TELEFON"], iso_date(row["Sıradaki Test"]), row["Durum"]
        ) for row in results],
    )

    for filename in ("BI-Psiko 2025.xlsx", "BI-Psiko 2026.xlsx"):
        pool = rows_from_workbook(folder / filename)
        connection.executemany(
            "INSERT INTO raw_external_pool VALUES (?, ?, ?, ?, ?)",
            [(
                row["SN"], row["Customer"], iso_date(row["TARİH"]), row["TELEFON"], filename
            ) for row in pool],
        )


def build_connection():
    connection = sqlite3.connect(":memory:")
    connection.row_factory = sqlite3.Row
    connection.executescript((SQL_DIR / SQL_FILES[0]).read_text(encoding="utf-8"))
    with TemporaryDirectory() as directory:
        folder = Path(directory)
        load_setup_module().build_data(folder)
        load_workbooks(connection, folder)
    for filename in SQL_FILES[1:]:
        connection.executescript((SQL_DIR / filename).read_text(encoding="utf-8"))
    return connection


def markdown_table(rows):
    if not rows:
        return "_No rows._"
    columns = rows[0].keys()
    lines = ["| " + " | ".join(columns) + " |", "|" + "---|" * len(columns)]
    for row in rows:
        values = ["" if row[column] is None else str(row[column]) for column in columns]
        lines.append("| " + " | ".join(values) + " |")
    return "\n".join(lines)


def render_results(connection):
    sections = [
        "# Executed SQL results",
        "",
        "> Generated from the same fully synthetic Excel workbooks used by the public Power BI demo. These are portfolio-demo outputs, not private or operational results.",
        "",
        "**As-of date:** 2026-09-08; legal anchor 2021-06-30; five-year renewal cycles. Four future-dated external demo rows are retained in raw data but excluded from as-of marts.",
        "",
        "**Scope note:** the synthetic contact KPI below does not reproduce or reconcile the reported 648 / 1,092 summary with the screenshot's 211 / 396 result. Their cohort, date and filter relationship remains undocumented.",
    ]
    for title, query in REPORTS:
        sections.extend(["", f"## {title}", "", markdown_table(connection.execute(query).fetchall())])
    return "\n".join(sections) + "\n"


def main():
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--write-results", action="store_true",
                        help="Refresh sql/RESULTS.md from the executed queries")
    args = parser.parse_args()
    connection = build_connection()
    try:
        output = render_results(connection)
        print(output, end="")
        failures = connection.execute(
            "SELECT check_name, issue_count FROM quality_check_results WHERE issue_count <> 0"
        ).fetchall()
        if failures:
            raise SystemExit("SQL demo failed one or more data-quality checks.")
        if args.write_results:
            (SQL_DIR / "RESULTS.md").write_text(output, encoding="utf-8")
            print("Updated sql/RESULTS.md.")
        print("PASS: all SQL data-quality checks returned zero issues.")
    finally:
        connection.close()


if __name__ == "__main__":
    main()
